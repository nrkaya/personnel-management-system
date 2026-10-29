"""Security, privacy and role-separation regression tests."""
import glob
import io
import re
import shutil
import tempfile
import unittest
from datetime import date, timedelta

from openpyxl import load_workbook
from PIL import Image
from werkzeug.security import generate_password_hash

import config
from app import create_app
from database import db
from models import Izin, Kullanici, Log, LoginAttempt, Personel
from utils import utcnow

PASSWORD = 'Str0ng!Passw0rd'


class SecurityTestCase(unittest.TestCase):
    def setUp(self):
        self.app = create_app('testing')
        self.upload_dir = tempfile.mkdtemp()
        self.legacy_dir = tempfile.mkdtemp()
        self.app.config.update(UPLOAD_FOLDER=self.upload_dir, LEGACY_UPLOAD_FOLDER=self.legacy_dir)
        self.client = self.app.test_client()
        self.app_context = self.app.app_context()
        self.app_context.push()
        db.create_all()

        self.admin = self._user('admin', 'Admin')
        self.ik = self._user('ik', 'IK')
        self.mudur = self._user('mudur', 'Mudur', departman='Satış')
        self.muhasebe = self._user('muhasebe', 'Muhasebe')
        self.sales = Personel(ad='Satis', soyad='Kisi', departman='Satış', maas=1000,
                              dogum_tarihi=date(1990, 1, 2), ek_bilgi='GIZLI-NOT')
        self.it = Personel(ad='Bilgi', soyad='Islem', departman='IT', maas=2000)
        db.session.add_all([self.sales, self.it])
        db.session.commit()

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.app_context.pop()
        shutil.rmtree(self.upload_dir, ignore_errors=True)
        shutil.rmtree(self.legacy_dir, ignore_errors=True)

    def _user(self, name, rol, departman=None):
        user = Kullanici(kullanici_adi=name, sifre_hash=generate_password_hash(PASSWORD),
                         rol=rol, departman=departman, aktif_mi=True)
        db.session.add(user)
        db.session.commit()
        return user

    def _login(self, user, client=None):
        client = client or self.client
        return client.post('/login', data={'kullanici_adi': user.kullanici_adi, 'sifre': PASSWORD})

    def _text(self, response):
        return response.get_data(as_text=True)

    # --- Authentication, lockout and sessions -------------------------------

    def test_other_login_does_not_reset_lockout_counter(self):
        for _ in range(4):
            self.client.post('/login', data={'kullanici_adi': 'admin', 'sifre': 'wrong'})
        self._login(self.muhasebe)
        self.client.post('/logout')
        attempt = LoginAttempt.query.filter_by(kullanici_adi='admin').first()
        self.assertIsNotNone(attempt)
        self.assertEqual(attempt.deneme_sayisi, 4)

        self.client.post('/login', data={'kullanici_adi': 'admin', 'sifre': 'wrong'})
        locked = self._login(self.admin)
        self.assertIn('kilitlendi', self._text(locked))
        self.assertEqual(self.client.get('/users').status_code, 302)

    def test_older_locked_row_still_blocks_login(self):
        now = utcnow()
        db.session.add(LoginAttempt(kullanici_adi='admin', ip_adresi='10.0.0.9', deneme_sayisi=5,
                                    son_deneme=now - timedelta(minutes=1), kilitli_mi=True,
                                    kilit_bitis=now + timedelta(minutes=9)))
        # A newer, unlocked row for the same account must not hide the lock.
        db.session.add(LoginAttempt(kullanici_adi='admin', ip_adresi='10.0.0.8', deneme_sayisi=1,
                                    son_deneme=now, kilitli_mi=False))
        db.session.commit()
        response = self._login(self.admin)
        self.assertIn('kilitlendi', self._text(response))

    def test_failure_message_is_the_same_for_unknown_users(self):
        known = self._text(self.client.post('/login', data={'kullanici_adi': 'admin', 'sifre': 'x'}))
        unknown = self._text(self.client.post('/login', data={'kullanici_adi': 'nobody', 'sifre': 'x'}))
        self.assertIn('Geçersiz kullanıcı adı veya şifre', known)
        self.assertIn('Geçersiz kullanıcı adı veya şifre', unknown)
        self.assertNotIn('Kalan deneme', known + unknown)

    def test_typed_unknown_username_is_not_stored(self):
        self.client.post('/login', data={'kullanici_adi': 'MySecretP@ss1', 'sifre': 'x'})
        self.assertEqual(Log.query.filter(Log.kullanici_adi.contains('MySecretP@ss1')).count(), 0)
        self.assertEqual(LoginAttempt.query.filter(LoginAttempt.kullanici_adi.contains('MySecretP@ss1')).count(), 0)

    def test_logout_revokes_copied_session_cookie(self):
        self._login(self.admin)
        cookie = self.client.get_cookie('session').value
        self.client.post('/logout')
        replay = self.app.test_client()
        replay.set_cookie('session', cookie)
        self.assertEqual(replay.get('/users').status_code, 302)

    def test_password_change_keeps_current_session_and_revokes_others(self):
        other_device = self.app.test_client()
        self._login(self.ik, other_device)
        self._login(self.ik)
        response = self.client.post('/hesap/sifre', data={
            'mevcut_sifre': PASSWORD, 'yeni_sifre': 'An0ther!Passw0rd', 'yeni_sifre_tekrar': 'An0ther!Passw0rd',
        })
        self.assertEqual(response.status_code, 302)
        self.assertEqual(self.client.get('/personeller').status_code, 200)
        self.assertEqual(other_device.get('/personeller').status_code, 302)

    def test_password_change_requires_current_password(self):
        self._login(self.ik)
        old_hash = self.ik.sifre_hash
        self.client.post('/hesap/sifre', data={
            'mevcut_sifre': 'wrong', 'yeni_sifre': 'An0ther!Passw0rd', 'yeni_sifre_tekrar': 'An0ther!Passw0rd',
        })
        db.session.refresh(self.ik)
        self.assertEqual(self.ik.sifre_hash, old_hash)

    def test_leaked_secrets_are_rejected_in_production(self):
        original = config.ProductionConfig.SECRET_KEY
        config.ProductionConfig.SECRET_KEY = 'kaya-pts-ultra-secure-key-2026-v4'
        try:
            with self.assertRaises(RuntimeError):
                create_app('production')
        finally:
            config.ProductionConfig.SECRET_KEY = original

    def test_csrf_tokens_live_as_long_as_the_session(self):
        self.assertEqual(
            self.app.config['WTF_CSRF_TIME_LIMIT'],
            int(self.app.config['PERMANENT_SESSION_LIFETIME'].total_seconds()),
        )

    def test_language_switch_rejects_backslash_redirect(self):
        for target in ('/%5Cevil.example', '//evil.example/path', 'https:evil.example'):
            response = self.client.get(f'/lang/tr?next={target}')
            self.assertEqual(response.headers['Location'], '/login', target)
        response = self.client.get('/lang/en?next=/personeller%3Fpage%3D2')
        self.assertEqual(response.headers['Location'], '/personeller?page=2')

    # --- Role separation ------------------------------------------------------

    def test_dashboard_audit_trail_is_admin_only(self):
        self._login(self.muhasebe)
        self.assertNotIn('Sisteme başarıyla giriş yaptı', self._text(self.client.get('/dashboard')))
        self.client.post('/logout')
        self._login(self.admin)
        self.assertIn('Sisteme başarıyla giriş yaptı', self._text(self.client.get('/dashboard')))

    def test_ik_cannot_hard_delete_and_sees_no_button(self):
        self.sales.aktif_mi = False
        db.session.commit()
        self._login(self.ik)
        self.assertNotIn('hard_delete', self._text(self.client.get('/trash')))
        self.assertEqual(self.client.post(f'/personel/hard_delete/{self.sales.id}').status_code, 403)

    def test_manager_cannot_delete_leaves_and_sees_no_button(self):
        izin = Izin(personel_id=self.sales.id, baslangic=date.today(), bitis=date.today(), tur='Yıllık İzin')
        db.session.add(izin)
        db.session.commit()
        self._login(self.mudur)
        self.assertNotIn('/izin/sil/', self._text(self.client.get(f'/personel/{self.sales.id}')))
        self.assertEqual(self.client.post(f'/izin/sil/{izin.id}').status_code, 403)

    def test_sensitive_fields_hidden_from_accounting(self):
        db.session.add(Izin(personel_id=self.sales.id, baslangic=date.today(), bitis=date.today(),
                            tur='Hastalık Raporu', aciklama='SAGLIK-BILGISI'))
        db.session.commit()
        self._login(self.muhasebe)
        profile = self._text(self.client.get(f'/personel/{self.sales.id}'))
        self.assertNotIn('GIZLI-NOT', profile)
        self.assertNotIn('02.01.1990', profile)
        self.assertNotIn('SAGLIK-BILGISI', profile)
        self.assertNotIn('SAGLIK-BILGISI', self._text(self.client.get('/izinler')))

    def test_manager_sees_leave_reason_hr_sees_personal_data(self):
        db.session.add(Izin(personel_id=self.sales.id, baslangic=date.today(), bitis=date.today(),
                            tur='Mazeret İzni', aciklama='SEBEP-METNI'))
        db.session.commit()
        self._login(self.mudur)
        self.assertIn('SEBEP-METNI', self._text(self.client.get('/izinler')))
        self.client.post('/logout')
        self._login(self.ik)
        self.assertIn('GIZLI-NOT', self._text(self.client.get(f'/personel/{self.sales.id}')))

    # --- A: department scope for managers -------------------------------------

    def test_manager_only_sees_own_department(self):
        self._login(self.mudur)
        listing = self._text(self.client.get('/personeller'))
        self.assertIn('Satis Kisi', listing)
        self.assertNotIn('Bilgi Islem', listing)
        self.assertEqual(self.client.get(f'/personel/{self.it.id}').status_code, 404)
        self.assertEqual(self.client.post(f'/izin/ekle/{self.it.id}', data={
            'baslangic': '2026-01-01', 'bitis': '2026-01-02', 'tur': 'Yıllık İzin',
        }).status_code, 404)
        export = load_workbook(io.BytesIO(self.client.get('/export').data)).active
        names = [row[1].value for row in export.iter_rows(min_row=2)]
        self.assertEqual(names, ['Satis'])

    def test_manager_cannot_change_leave_of_other_department(self):
        izin = Izin(personel_id=self.it.id, baslangic=date.today(), bitis=date.today(), tur='Yıllık İzin')
        db.session.add(izin)
        db.session.commit()
        self._login(self.mudur)
        self.assertEqual(self.client.post(f'/izin/durum/{izin.id}', data={'durum': 'Onaylandı'}).status_code, 404)
        self.assertEqual(db.session.get(Izin, izin.id).durum, 'Beklemede')

    def test_manager_without_department_sees_nobody(self):
        self.mudur.departman = None
        db.session.commit()
        self._login(self.mudur)
        listing = self._text(self.client.get('/personeller'))
        self.assertNotIn('Satis Kisi', listing)
        self.assertIn('departman atanmadığı', listing)

    def test_manager_role_requires_department(self):
        self._login(self.admin)
        self.client.post('/users/add', data={
            'kullanici_adi': 'yenimudur', 'sifre': PASSWORD, 'rol': 'Mudur', 'departman': '',
        })
        self.assertIsNone(Kullanici.query.filter_by(kullanici_adi='yenimudur').first())

    # --- User management -------------------------------------------------------

    def test_admin_reset_password_revokes_sessions_and_lock(self):
        target = self.app.test_client()
        self._login(self.ik, target)
        db.session.add(LoginAttempt(kullanici_adi='ik', ip_adresi='1.2.3.4', deneme_sayisi=5, kilitli_mi=True,
                                    kilit_bitis=utcnow() + timedelta(minutes=5)))
        db.session.commit()
        self._login(self.admin)
        self.client.post(f'/users/reset_password/{self.ik.id}', data={
            'sifre': 'Fresh!Passw0rd1', 'sifre_tekrar': 'Fresh!Passw0rd1',
        })
        self.assertEqual(target.get('/personeller').status_code, 302)
        self.assertEqual(LoginAttempt.query.filter_by(kullanici_adi='ik').count(), 0)

    def test_deactivating_user_revokes_access(self):
        target = self.app.test_client()
        self._login(self.muhasebe, target)
        self._login(self.admin)
        self.client.post(f'/users/edit/{self.muhasebe.id}', data={'rol': 'Muhasebe'})
        db.session.refresh(self.muhasebe)
        self.assertFalse(self.muhasebe.aktif_mi)
        self.assertEqual(target.get('/personeller').status_code, 302)

    def test_admin_cannot_demote_or_disable_self(self):
        self._login(self.admin)
        self.client.post(f'/users/edit/{self.admin.id}', data={'rol': 'Muhasebe'})
        db.session.refresh(self.admin)
        self.assertEqual(self.admin.rol, 'Admin')
        self.assertTrue(self.admin.aktif_mi)

    # --- Personnel data ----------------------------------------------------------

    def test_excel_export_neutralises_formulas(self):
        db.session.add(Personel(ad='=HYPERLINK("http://evil.example","x")', soyad='+cmd', maas=1))
        db.session.commit()
        self._login(self.admin)
        sheet = load_workbook(io.BytesIO(self.client.get('/export').data)).active
        for row in sheet.iter_rows():
            for cell in row:
                self.assertNotEqual(cell.data_type, 'f', cell.value)

    def test_edit_with_missing_fields_keeps_values(self):
        self._login(self.ik)
        response = self.client.post(f'/edit/{self.it.id}', data={'pozisyon': 'Uzman'})
        self.assertEqual(response.status_code, 302)
        db.session.refresh(self.it)
        self.assertEqual(self.it.pozisyon, 'Uzman')
        self.assertEqual(self.it.ad, 'Bilgi')
        self.assertEqual(self.it.departman, 'IT')

    def test_invalid_personnel_values_are_rejected(self):
        self._login(self.ik)
        for data in ({'departman': 'x' * 101}, {'telefon': '<script>'}, {'maas': '99999999999'},
                     {'dogum_tarihi': (date.today() + timedelta(days=1)).isoformat()}):
            self.client.post(f'/edit/{self.it.id}', data=data)
        db.session.refresh(self.it)
        self.assertEqual(self.it.departman, 'IT')
        self.assertIsNone(self.it.telefon)
        self.assertEqual(int(self.it.maas), 2000)
        self.assertIsNone(self.it.dogum_tarihi)

    def test_photos_are_private_and_randomly_named(self):
        image = io.BytesIO()
        Image.new('RGB', (4, 4), 'red').save(image, 'PNG')
        image.seek(0)
        self._login(self.ik)
        self.client.post(f'/edit/{self.sales.id}', data={'avatar': (image, 'foto.png')},
                         content_type='multipart/form-data')
        db.session.refresh(self.sales)
        filename = self.sales.avatar_path
        self.assertRegex(filename, r'^[0-9a-f]{32}\.png$')
        self.assertEqual(self.client.get(f'/static/uploads/{filename}').status_code, 404)
        self.assertEqual(self._photo_status(self.client, self.sales), 200)

        anonymous = self.app.test_client()
        self.assertEqual(self._photo_status(anonymous, self.sales), 302)

        self.client.post('/logout')
        self._login(self.mudur)
        self.assertEqual(self._photo_status(self.client, self.sales), 200)
        self.assertEqual(self._photo_status(self.client, self.it), 404)

    def _photo_status(self, client, person):
        response = client.get(f'/personel/{person.id}/foto')
        response.close()
        return response.status_code

    def test_leave_reason_length_and_pagination(self):
        self._login(self.ik)
        self.client.post(f'/izin/ekle/{self.sales.id}', data={
            'baslangic': '2026-01-01', 'bitis': '2026-01-02', 'tur': 'Yıllık İzin', 'aciklama': 'x' * 1001,
        })
        self.assertEqual(Izin.query.count(), 0)
        for day in range(35):
            start = date(2025, 1, 1) + timedelta(days=day * 2)
            db.session.add(Izin(personel_id=self.it.id, baslangic=start, bitis=start, tur='Yıllık İzin'))
        db.session.commit()
        self.assertIn('1 / 2', self._text(self.client.get('/izinler')))

    # --- API scopes ----------------------------------------------------------------

    def test_api_salary_and_personal_fields_need_scopes(self):
        self.app.config['API_KEYS_JSON'] = '{"directory": {"key": "directory-key-0123456789abcdef", "scopes": ["personnel:read", "personnel:write"]}}'
        headers = {'X-API-KEY': 'directory-key-0123456789abcdef'}
        person = self.client.get(f'/api/v1/personel/{self.sales.id}', headers=headers).get_json()['data']
        self.assertNotIn('maas', person)
        self.assertNotIn('dogum_tarihi', person)
        self.assertNotIn('ek_bilgi', person)
        self.assertEqual(self.client.put(f'/api/v1/personel/{self.sales.id}', json={'maas': 1}, headers=headers).status_code, 403)

        full = self.client.get(f'/api/v1/personel/{self.sales.id}', headers={'X-API-KEY': self.app.config['API_KEY']})
        self.assertIn('maas', full.get_json()['data'])

    # --- Front end and maintenance ------------------------------------------------

    def test_csp_forbids_inline_scripts(self):
        csp = self.client.get('/login').headers['Content-Security-Policy']
        self.assertNotIn('unsafe-inline', csp)
        self.assertNotIn('https://', csp)

    def test_templates_have_no_inline_handlers_or_scripts(self):
        pattern = re.compile(r'\son[a-z]+="|<script>')
        for path in glob.glob('templates/*.html'):
            with open(path, encoding='utf-8') as handle:
                self.assertIsNone(pattern.search(handle.read()), path)

    def test_purge_logs_applies_retention(self):
        db.session.add(Log(kullanici_adi='eski', islem_tipi='Giriş', tarih_saat=utcnow() - timedelta(days=400)))
        db.session.add(Log(kullanici_adi='yeni', islem_tipi='Giriş', tarih_saat=utcnow()))
        db.session.commit()
        result = self.app.test_cli_runner().invoke(args=['purge-logs', '--days', '365'])
        self.assertEqual(result.exit_code, 0, result.output)
        self.assertEqual(Log.query.filter_by(kullanici_adi='eski').count(), 0)
        self.assertEqual(Log.query.filter_by(kullanici_adi='yeni').count(), 1)


if __name__ == '__main__':
    unittest.main()
