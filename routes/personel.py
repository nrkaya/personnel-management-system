import io
import math
from decimal import Decimal, InvalidOperation
from datetime import datetime, date
from flask import Blueprint, render_template, request, redirect, url_for, flash, send_file, abort
from sqlalchemy import or_
from sqlalchemy.exc import SQLAlchemyError
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from database import db
from i18n import _
from models import Personel, Izin
from utils import (
    DEFAULT_AVATAR_STATIC, MAX_SALARY, PHONE_PATTERN, avatar_file_path, can_access_personnel,
    get_active_session_user, has_permission, log_action, permission_required, remove_avatar,
    save_avatar, scope_personnel, validate_email,
)

personel_bp = Blueprint('personel', __name__)

# Form field -> (maximum length, translation key of its label)
TEXT_FIELDS = {
    'ad': (100, 'modal_first_name'),
    'soyad': (100, 'modal_last_name'),
    'departman': (100, 'modal_dept'),
    'pozisyon': (100, 'modal_pos'),
    'telefon': (30, 'modal_phone'),
    'email': (120, 'modal_email'),
    'ek_bilgi': (10_000, 'modal_notes'),
}
# Spreadsheet apps evaluate text starting with these characters as formulas.
FORMULA_PREFIXES = ('=', '+', '-', '@', '\t', '\r')


def _parse_form_date(value, label_key):
    if not value:
        return None
    try:
        return datetime.strptime(value, '%Y-%m-%d').date()
    except ValueError as exc:
        raise ValueError(_('msg_invalid_date', field=_(label_key))) from exc


def _parse_salary(value):
    if not value:
        return Decimal('0.00')
    try:
        salary = Decimal(value).quantize(Decimal('0.01'))
    except InvalidOperation as exc:
        raise ValueError(_('msg_invalid_salary')) from exc
    if not salary.is_finite() or salary < 0 or salary > MAX_SALARY:
        raise ValueError(_('msg_invalid_salary'))
    return salary


# Fields that need the 'personal' permission to set.
PERSONAL_FIELDS = {'ek_bilgi', 'dogum_tarihi'}


def _clean_personnel_form(form, personel=None, can_set_salary=False, can_set_personal=False):
    """Validate the add/edit form and return the values to store.

    Nothing on the model is touched until every field is valid. When editing,
    a field missing from the submitted form keeps its current value; fields
    the user's role may not set are ignored.
    """
    is_new = personel is None
    data = {}
    for field, (max_length, label_key) in TEXT_FIELDS.items():
        if field in PERSONAL_FIELDS and not can_set_personal:
            continue
        raw = form.get(field)
        if raw is None and not is_new:
            continue
        value = (raw or '').strip()
        if len(value) > max_length:
            raise ValueError(_('msg_field_too_long', field=_(label_key), max=max_length))
        data[field] = value or None

    if ('ad' in data and not data['ad']) or ('soyad' in data and not data['soyad']):
        raise ValueError(_('msg_name_required'))
    if data.get('email') and not validate_email(data['email']):
        raise ValueError(_('msg_invalid_email'))
    if data.get('telefon') and not PHONE_PATTERN.fullmatch(data['telefon']):
        raise ValueError(_('msg_invalid_phone'))

    if can_set_salary and (is_new or form.get('maas') is not None):
        data['maas'] = _parse_salary((form.get('maas') or '').strip())

    if is_new or form.get('ise_giris_tarihi') is not None:
        ise_giris = _parse_form_date(form.get('ise_giris_tarihi'), 'modal_start_date')
        data['ise_giris_tarihi'] = ise_giris or (date.today() if is_new else None)
    if can_set_personal and (is_new or form.get('dogum_tarihi') is not None):
        dogum_tarihi = _parse_form_date(form.get('dogum_tarihi'), 'modal_birth_date')
        if dogum_tarihi and dogum_tarihi > date.today():
            raise ValueError(_('msg_birth_date_future'))
        data['dogum_tarihi'] = dogum_tarihi
    return data


def _get_visible_personnel(id, include_inactive=False):
    """Fetch a record the current user may see; anything else is a 404."""
    user = get_active_session_user()
    p = db.get_or_404(Personel, id)
    if not can_access_personnel(p, user):
        abort(404)
    if not p.aktif_mi and not (include_inactive and has_permission(user.rol, 'trash')):
        abort(404)
    return p


@personel_bp.route('/personeller')
@permission_required('view')
def index():
    user = get_active_session_user()
    search_query = request.args.get('q', '').strip()
    filter_dept = request.args.get('departman', '').strip()
    filter_pos = request.args.get('pozisyon', '').strip()
    page = request.args.get('page', 1, type=int)
    if page < 1:
        page = 1
    per_page = 20

    query = scope_personnel(Personel.query.filter_by(aktif_mi=True), user)

    if search_query:
        search_terms = f"%{search_query}%"
        query = query.filter(or_(
            Personel.ad.ilike(search_terms),
            Personel.soyad.ilike(search_terms),
            Personel.telefon.ilike(search_terms),
            Personel.email.ilike(search_terms)
        ))

    if filter_dept:
        query = query.filter(Personel.departman == filter_dept)

    if filter_pos:
        query = query.filter(Personel.pozisyon == filter_pos)

    total_items = query.count()
    total_pages = math.ceil(total_items / per_page) if total_items > 0 else 1
    if page > total_pages:
        page = total_pages

    personeller = query.order_by(Personel.id.desc()).offset((page - 1) * per_page).limit(per_page).all()

    departments = [r[0] for r in scope_personnel(db.session.query(Personel.departman).filter(
        Personel.aktif_mi == True, Personel.departman.isnot(None), Personel.departman != ''
    ), user).distinct().all()]

    positions = [r[0] for r in scope_personnel(db.session.query(Personel.pozisyon).filter(
        Personel.aktif_mi == True, Personel.pozisyon.isnot(None), Personel.pozisyon != ''
    ), user).distinct().all()]

    return render_template(
        'index.html',
        personeller=personeller,
        departments=departments,
        positions=positions,
        search_query=search_query,
        filter_dept=filter_dept,
        filter_pos=filter_pos,
        page=page,
        total_pages=total_pages,
        total_items=total_items,
        per_page=per_page,
        can_edit=has_permission(user.rol, 'edit'),
        can_delete=has_permission(user.rol, 'delete'),
        can_add=has_permission(user.rol, 'add'),
        can_view_salary=has_permission(user.rol, 'salary')
    )

@personel_bp.route('/personel/<int:id>')
@permission_required('view')
def profile(id):
    user = get_active_session_user()
    p = _get_visible_personnel(id, include_inactive=True)
    izinler = Izin.query.filter_by(personel_id=id).order_by(Izin.baslangic.desc()).all()

    return render_template(
        'profile.html',
        p=p,
        izinler=izinler,
        can_edit=has_permission(user.rol, 'edit'),
        can_delete=has_permission(user.rol, 'delete'),
        can_view_salary=has_permission(user.rol, 'salary'),
        can_manage_leaves=has_permission(user.rol, 'leaves') and p.aktif_mi
    )


@personel_bp.route('/personel/<int:id>/foto')
@permission_required('view')
def avatar(id):
    """Serve a personnel photo to logged-in users who may see that record."""
    p = _get_visible_personnel(id, include_inactive=True)
    path = avatar_file_path(p.avatar_filename)
    if path is None:
        return redirect(url_for('static', filename=DEFAULT_AVATAR_STATIC))
    response = send_file(path, conditional=True)
    response.headers['Cache-Control'] = 'private, max-age=3600'
    return response

@personel_bp.route('/add', methods=['POST'])
@permission_required('add')
def add():
    user = get_active_session_user()
    try:
        data = _clean_personnel_form(
            request.form,
            can_set_salary=has_permission(user.rol, 'salary'),
            can_set_personal=has_permission(user.rol, 'personal'),
        )
    except ValueError as exc:
        flash(str(exc), 'error')
        return redirect(url_for('personel.index'))

    # Handle avatar file upload
    avatar_file = request.files.get('avatar')
    avatar_path = 'default.svg'
    if avatar_file and avatar_file.filename != '':
        success, msg, filename = save_avatar(avatar_file)
        if success:
            avatar_path = filename
        else:
            flash(_('msg_photo_upload_failed', reason=msg), 'error')
            return redirect(url_for('personel.index'))

    new_personel = Personel(**data, avatar_path=avatar_path, aktif_mi=True)
    try:
        db.session.add(new_personel)
        db.session.commit()
    except SQLAlchemyError:
        db.session.rollback()
        remove_avatar(avatar_path)
        flash(_('msg_personnel_create_failed'), 'error')
        return redirect(url_for('personel.index'))

    log_action(user.kullanici_adi, 'Personel Eklendi', f"{new_personel.tam_ad} (ID: {new_personel.id}) sisteme eklendi.")
    flash(_('msg_personnel_added', name=new_personel.tam_ad), 'success')
    return redirect(url_for('personel.index'))

@personel_bp.route('/edit/<int:id>', methods=['POST'])
@permission_required('edit')
def edit(id):
    user = get_active_session_user()
    p = db.get_or_404(Personel, id)
    try:
        data = _clean_personnel_form(
            request.form, p,
            can_set_salary=has_permission(user.rol, 'salary'),
            can_set_personal=has_permission(user.rol, 'personal'),
        )
    except ValueError as exc:
        flash(str(exc), 'error')
        return redirect(url_for('personel.profile', id=id))

    old_avatar_path = p.avatar_path
    new_avatar_path = None
    avatar_file = request.files.get('avatar')
    if avatar_file and avatar_file.filename != '':
        success, msg, filename = save_avatar(avatar_file)
        if not success:
            flash(_('msg_photo_update_failed', reason=msg), 'error')
            return redirect(url_for('personel.profile', id=id))
        new_avatar_path = filename

    for field, value in data.items():
        setattr(p, field, value)
    if new_avatar_path:
        p.avatar_path = new_avatar_path

    try:
        db.session.commit()
    except SQLAlchemyError:
        db.session.rollback()
        if new_avatar_path:
            remove_avatar(new_avatar_path)
        flash(_('msg_personnel_update_failed'), 'error')
        return redirect(url_for('personel.profile', id=id))
    if new_avatar_path:
        remove_avatar(old_avatar_path)
    log_action(user.kullanici_adi, 'Personel Güncellendi', f"{p.tam_ad} (ID: {p.id}) bilgileri güncellendi.")
    flash(_('msg_personnel_updated', name=p.tam_ad), 'success')
    return redirect(url_for('personel.profile', id=id))

@personel_bp.route('/delete/<int:id>', methods=['POST'])
@permission_required('delete')
def delete(id):
    user = get_active_session_user()
    p = db.get_or_404(Personel, id)
    if not p.aktif_mi:
        flash(_('msg_personnel_already_in_trash'), 'info')
        return redirect(url_for('personel.index'))
    p.aktif_mi = False
    db.session.commit()

    log_action(user.kullanici_adi, 'Personel Silindi (Çöp Kutusu)', f"{p.tam_ad} (ID: {p.id}) arşive taşındı.")
    flash(_('msg_personnel_trashed', name=p.tam_ad), 'success')
    return redirect(url_for('personel.index'))

@personel_bp.route('/trash')
@permission_required('trash')
def trash():
    deleted_personnel = Personel.query.filter_by(aktif_mi=False).order_by(Personel.guncellenme_tarihi.desc()).all()
    return render_template('trash.html', personeller=deleted_personnel)

@personel_bp.route('/personel/restore/<int:id>', methods=['POST'])
@permission_required('trash')
def restore(id):
    user = get_active_session_user()
    p = db.get_or_404(Personel, id)
    if p.aktif_mi:
        flash(_('msg_personnel_already_active'), 'error')
        return redirect(url_for('personel.trash'))
    p.aktif_mi = True
    db.session.commit()

    log_action(user.kullanici_adi, 'Personel Geri Yüklendi', f"{p.tam_ad} (ID: {p.id}) arşivden geri yüklendi.")
    flash(_('msg_personnel_restored', name=p.tam_ad), 'success')
    return redirect(url_for('personel.trash'))

@personel_bp.route('/personel/hard_delete/<int:id>', methods=['POST'])
@permission_required('hard_delete')
def hard_delete(id):
    user = get_active_session_user()
    p = db.get_or_404(Personel, id)
    if p.aktif_mi:
        abort(400, description='Yalnızca çöp kutusundaki personel kalıcı olarak silinebilir.')
    tam_ad = p.tam_ad
    avatar_path = p.avatar_path
    db.session.delete(p)
    db.session.commit()
    remove_avatar(avatar_path)

    log_action(user.kullanici_adi, 'Kalıcı Silme', f"Personel (ID: {id}) sistemden kalıcı olarak silindi.")
    flash(_('msg_personnel_hard_deleted', name=tam_ad), 'success')
    return redirect(url_for('personel.trash'))


def _mark_as_text(cell):
    """Store a value as plain text so spreadsheet apps never run it as a formula."""
    if isinstance(cell.value, str) and cell.value.startswith(FORMULA_PREFIXES):
        cell.data_type = 's'
        cell.quotePrefix = True


@personel_bp.route('/export')
@permission_required('export')
def export_excel():
    """Export formatted personnel dataset to Excel (.xlsx)."""
    user = get_active_session_user()
    can_see_salary = has_permission(user.rol, 'salary')

    personeller = scope_personnel(Personel.query.filter_by(aktif_mi=True), user).order_by(Personel.id.asc()).all()

    wb = Workbook()
    ws = wb.active
    # Sheet titles may not exceed 31 characters.
    ws.title = _('export_sheet_title')[:31]

    headers = ["ID"] + [_(key) for key in (
        'modal_first_name', 'modal_last_name', 'modal_dept', 'modal_pos',
        'modal_phone', 'modal_email', 'modal_start_date',
    )]
    if can_see_salary:
        headers.append(_('modal_salary'))

    ws.append(headers)

    header_font = Font(name='Segoe UI', size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2563EB")
    alignment_center = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = alignment_center
        cell.border = thin_border

    for row_idx, p in enumerate(personeller, start=2):
        row_data = [
            p.id,
            p.ad,
            p.soyad,
            p.departman or '-',
            p.pozisyon or '-',
            p.telefon or '-',
            p.email or '-',
            p.ise_giris_tarihi.strftime('%Y-%m-%d') if p.ise_giris_tarihi else '-'
        ]
        if can_see_salary:
            row_data.append(p.maas)

        ws.append(row_data)

        # Zebra striping
        fill_color = "F9FAFB" if row_idx % 2 == 0 else "FFFFFF"
        for col_idx in range(1, len(row_data) + 1):
            c = ws.cell(row=row_idx, column=col_idx)
            _mark_as_text(c)
            c.fill = PatternFill("solid", fgColor=fill_color)
            c.border = thin_border
            if col_idx in [1, 6, 8]:
                c.alignment = alignment_center

    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_length:
                max_length = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_length + 4, 12)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    log_action(user.kullanici_adi, 'Dışa Aktarma', f'Personel listesi Excel (.xlsx) olarak indirildi ({len(personeller)} kayıt).')
    return send_file(
        out,
        download_name=f"KAYA_PTS_{_('export_file_name')}_{datetime.now().strftime('%Y%m%d')}.xlsx",
        as_attachment=True,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
